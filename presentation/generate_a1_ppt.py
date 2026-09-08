from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


import os, sys

# Output path: first CLI argument, else next to this script.
_HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    _HERE, "Project_A1_Minimal_Format_Presentation.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(18, 35, 56)
BLUE = RGBColor(42, 114, 166)
TEAL = RGBColor(39, 151, 135)
INK = RGBColor(36, 45, 55)
MUTED = RGBColor(95, 110, 125)
PALE = RGBColor(239, 245, 248)
WHITE = RGBColor(255, 255, 255)
SOFT_BLUE = RGBColor(224, 238, 248)
SOFT_TEAL = RGBColor(224, 244, 239)


def add_text(slide, x, y, w, h, value, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, items, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = INK
        paragraph.space_after = Pt(12)
        paragraph.bullet = True
    return box


def add_card(slide, x, y, w, h, title, body, fill=PALE, line=BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.0)
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.16)
    frame.margin_right = Inches(0.16)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p1 = frame.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = title
    r1.font.name = "Aptos"
    r1.font.size = Pt(18)
    r1.font.bold = True
    r1.font.color.rgb = NAVY
    p2 = frame.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = body
    r2.font.name = "Aptos"
    r2.font.size = Pt(13)
    r2.font.color.rgb = MUTED
    return shape


def base_slide(title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.16))
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()

    add_text(slide, 0.65, 0.48, 11.8, 0.55, title, 30, NAVY, True)
    if subtitle:
        add_text(slide, 0.68, 1.05, 11.6, 0.35, subtitle, 14, MUTED)
    return slide


# 1. Scanned guide-signed first slide
slide = base_slide("Scanned copy of the guide-signed first slide")
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.1), Inches(1.55), Inches(9.1), Inches(4.95))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(250, 252, 253)
shape.line.color.rgb = RGBColor(170, 185, 195)
shape.line.width = Pt(1.5)
add_text(
    slide,
    2.45,
    3.45,
    8.4,
    0.55,
    "Insert scanned guide-signed first slide here",
    24,
    MUTED,
    True,
    PP_ALIGN.CENTER,
)
add_text(
    slide,
    0.8,
    6.75,
    11.8,
    0.3,
    "Placeholder included because no signed scan image was found in the project folder.",
    12,
    MUTED,
    False,
    PP_ALIGN.CENTER,
)


# 2. Introduction
slide = base_slide("Introduction", "Project A1: big.LITTLE Heterogeneous Multi-Core SoC")
add_bullets(
    slide,
    0.9,
    1.75,
    11.6,
    3.8,
    [
        "Design a RISC-V based heterogeneous multi-core SoC using performance and efficiency cores.",
        "Use a big.LITTLE style architecture to balance compute performance and energy usage.",
        "Focus on RTL design, simulation and verification; FPGA implementation is not included.",
        "Optional extension: ASIC flow up to GDSII if time permits.",
    ],
)


# 3. Problem Statement
slide = base_slide("Problem Statement")
add_bullets(
    slide,
    0.9,
    1.75,
    11.6,
    3.4,
    [
        "Single-core designs cannot efficiently handle both high-performance and low-power workloads.",
        "A multi-core SoC requires correct communication between cores, memory and peripherals.",
        "Shared-memory systems need cache coherence to avoid stale or inconsistent data.",
        "The project addresses these challenges using a compact RISC-V heterogeneous SoC design.",
    ],
)


# 4. Objectives
slide = base_slide("Objectives")
add_bullets(
    slide,
    0.9,
    1.7,
    11.6,
    4.3,
    [
        "Design RV32I efficiency cores and RV32IM performance cores.",
        "Integrate cores through a simple shared-memory SoC architecture.",
        "Implement memory, interconnect and basic peripheral support.",
        "Verify processor functionality, memory access and system-level behavior.",
        "Generate synthesis reports and attempt ASIC physical design flow only if schedule allows.",
    ],
)


# 5. Proposed Methodology
slide = base_slide("Proposed Methodology")
add_card(slide, 0.75, 1.65, 2.25, 1.35, "Step 1", "Build and verify basic RV32I E-Core", SOFT_TEAL, TEAL)
add_card(slide, 3.25, 1.65, 2.25, 1.35, "Step 2", "Develop RV32IM P-Core with 5-stage pipeline", SOFT_BLUE, BLUE)
add_card(slide, 5.75, 1.65, 2.25, 1.35, "Step 3", "Add interconnect, SRAM, ROM, UART and timer", PALE, BLUE)
add_card(slide, 8.25, 1.65, 2.25, 1.35, "Step 4", "Integrate cores and run system tests", SOFT_TEAL, TEAL)
add_card(slide, 10.75, 1.65, 1.9, 1.35, "Step 5", "Synthesis and optional GDSII flow", PALE, BLUE)
add_bullets(
    slide,
    1.0,
    4.0,
    11.0,
    1.7,
    [
        "Develop each module with its own testbench before full SoC integration.",
        "Keep advanced cache coherence and migration features as stretch goals after the MVP is stable.",
    ],
    18,
)


# 6. Expected Outcomes
slide = base_slide("Expected Outcomes")
add_bullets(
    slide,
    0.9,
    1.7,
    11.6,
    4.3,
    [
        "A modular RTL design of a heterogeneous RISC-V multi-core SoC.",
        "Verified core execution for selected RISC-V instruction tests.",
        "Working SoC-level memory and peripheral communication.",
        "Simulation waveforms, logs and regression results as project evidence.",
        "Synthesis results and optional physical-design output if completed.",
    ],
)


# 7. Tools/Technologies to be Used
slide = base_slide("Tools/Technologies to be Used")
add_bullets(
    slide,
    0.9,
    1.65,
    11.6,
    4.6,
    [
        "SystemVerilog or Verilog for RTL design.",
        "Verilator / Icarus Verilog for simulation and linting.",
        "GTKWave for waveform analysis.",
        "RISC-V GCC toolchain for bare-metal test programs.",
        "Yosys for synthesis and area estimation.",
        "OpenROAD / OpenLane for optional RTL-to-GDSII exploration.",
        "Git and Makefiles for version control and repeatable runs.",
    ],
    18,
)


# 8. Project Timeline
slide = base_slide("Project Timeline")
timeline = [
    ("Weeks 1-2", "Architecture, ISA scope and interface definition"),
    ("Weeks 3-5", "E-Core RTL and unit verification"),
    ("Weeks 6-8", "P-Core RTL and pipeline verification"),
    ("Weeks 9-11", "Memory, interconnect and peripheral integration"),
    ("Weeks 12-14", "SoC-level verification and regression cleanup"),
    ("Weeks 15-16", "Synthesis, documentation and optional GDSII flow"),
]
for index, (time, work) in enumerate(timeline):
    y = 1.35 + index * 0.78
    add_card(slide, 0.95, y, 2.15, 0.5, time, "", SOFT_BLUE if index % 2 == 0 else SOFT_TEAL, BLUE)
    add_text(slide, 3.45, y + 0.09, 8.4, 0.25, work, 16, INK)


# 9. References
slide = base_slide("References")
add_bullets(
    slide,
    0.9,
    1.55,
    11.6,
    4.8,
    [
        "RISC-V Instruction Set Manual, Volume I: Unprivileged ISA.",
        "RISC-V Privileged Architecture Specification.",
        "ARM big.LITTLE processing architecture concept references.",
        "OpenROAD / OpenLane documentation for open-source ASIC flow.",
        "Project planning file: RISC-V_SoC_Project_Plan.md.",
        "Detailed notes file: Project_A1_Detailed_Design_Notes.md.",
    ],
    18,
)


prs.core_properties.title = "Project A1 - big.LITTLE Heterogeneous Multi-Core SoC"
prs.core_properties.subject = "Minimal project presentation"
prs.core_properties.author = ""
prs.save(OUT)
print(OUT)
